"""Write the engineering documents for a part record as Excel workbooks plus a markdown gate report.

One workbook per controlled document, named `<id>_<Document> (Rev - X).xlsx`. MOCK records get
MOCK in every file name and a banner on every sheet; ITAR records get the export-control warning.
"""

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .ap import action_priority
from .record import GATES, check, passed, trace
from .stackup import analyze

ITAR_WARNING = (
    "WARNING: This document contains technical data whose export is restricted by the Arms Export Control Act "
    "(22 U.S.C. 2751 et seq.) and the ITAR (22 CFR 120-130). Do not export or disclose to foreign persons without authorization."
)
MOCK_BANNER = "MOCK-UP ONLY. Built from invented or photo-scaled values. Not for bid, production or SAR use."

_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
_BANNER = Font(bold=True, color="C00000")
_AP_FILL = {"H": PatternFill("solid", fgColor="F4B6B6"), "M": PatternFill("solid", fgColor="FFE699"), "L": PatternFill("solid", fgColor="C6E0B4")}


def file_stem(record, doc):
    p = record["part"]
    pid = p["id"]
    if p.get("mock") and "MOCK" not in pid:
        pid += "_MOCK"
    return f"{pid}_{doc} (Rev - {p['rev']})"


def _sheet(wb, title, record, doc_title, columns, rows, widths=None):
    ws = wb.create_sheet(title[:31])
    p = record["part"]
    ws.append([doc_title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Part: {p['name']}", f"P/N: {p.get('pn', '')}", f"NSN: {p.get('nsn', '')}", f"Rev: {p['rev']}",
               f"Units: {p['units']}", f"Generated: {dt.date.today().isoformat()}"])
    if p.get("mock"):
        ws.append([MOCK_BANNER])
        ws.cell(ws.max_row, 1).font = _BANNER
    if p["export_control"].get("itar"):
        ws.append([ITAR_WARNING])
        ws.cell(ws.max_row, 1).font = _BANNER
    ws.append([])
    ws.append(columns)
    head = ws.max_row
    for i in range(1, len(columns) + 1):
        cell = ws.cell(head, i)
        cell.font, cell.fill = _HEAD, _HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(["" if v is None else v for v in r])
        for i in range(1, len(columns) + 1):
            ws.cell(ws.max_row, i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(head + 1, 1)
    for i, w in enumerate(widths or [18] * len(columns), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws, head


def _book(record, out_dir, doc, sheets):
    """sheets: list of (title, doc_title, columns, rows, widths, ap_col) where ap_col is the 1-based AP column or None."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, doc_title, columns, rows, widths, ap_col in sheets:
        ws, head = _sheet(wb, title, record, doc_title, columns, rows, widths)
        if ap_col:
            for r in range(head + 1, ws.max_row + 1):
                c = ws.cell(r, ap_col)
                if c.value in _AP_FILL:
                    c.fill = _AP_FILL[c.value]
    path = Path(out_dir) / f"{file_stem(record, doc)}.xlsx"
    wb.save(path)
    return path


def _tol_text(c):
    if c.get("requirement_text"):
        return c["requirement_text"]
    if "lsl" in c and "usl" in c:
        return f"{c['lsl']} - {c['usl']}"
    g = c.get("gdt", {})
    if g:
        mod = f" ({g['modifier']})" if g.get("modifier") else ""
        return f"{g.get('symbol', '')} {g.get('tol', '')}{mod} {'|'.join(g.get('datums', []))}".strip()
    return ""


def _fmea_rows(lines):
    rows = []
    for l in lines:
        rows.append([l["id"], l.get("op", ""), l["item"], l["function"], ", ".join(l.get("requirement_ids", []) + l.get("dfmea_ids", [])),
                     l["failure_mode"], l["effect"], l["S"], l["cause"], ", ".join(l.get("char_ids", [])),
                     l.get("prevention", ""), l["O"], l.get("detection", ""), l["D"],
                     action_priority(l["S"], l["O"], l["D"]), l.get("actions", ""), l.get("owner", ""),
                     l.get("due", ""), l.get("status", "")])
    return rows


_FMEA_COLS = ["ID", "Op", "Item / step", "Function", "Req / DFMEA", "Failure mode", "Failure effect", "S", "Failure cause",
              "Char", "Prevention controls", "O", "Detection controls", "D", "AP", "Recommended actions", "Owner", "Due", "Status"]
_FMEA_W = [6, 5, 20, 26, 7, 26, 28, 4, 28, 8, 26, 4, 26, 4, 5, 30, 10, 10, 10]


def build(record, out_dir):
    """Write every document the record has data for. Returns the list of files written."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    chars = record["characteristics"]
    by_id = {c["id"]: c for c in chars}
    files = []

    files.append(_book(record, out_dir, "Characteristics", [(
        "Characteristics", "Characteristics and balloon list (AS9102 Form 3 numbering)",
        ["Char", "Balloon", "Zone", "Feature", "Kind", "Requirement", "Nominal", "LSL", "USL", "Class", "Source", "Source note", "Req"],
        [[c["id"], c.get("balloon", ""), c.get("drawing_zone", ""), c["feature"], c["kind"], _tol_text(c), c.get("nominal"),
          c.get("lsl"), c.get("usl"), c["class"], c["source"], c.get("source_note", ""), ", ".join(c.get("requirement_ids", []))]
         for c in chars],
        [6, 7, 6, 30, 10, 26, 9, 9, 9, 6, 10, 34, 8], None)]))

    files.append(_book(record, out_dir, "Requirements", [(
        "Requirements", "Requirements",
        ["Req", "Type", "Requirement", "Source", "Characteristics"],
        [[q["id"], q["type"], q["text"], q["source"],
          ", ".join(c["id"] for c in chars if q["id"] in c.get("requirement_ids", []))] for q in record["requirements"]],
        [6, 12, 60, 24, 20], None)]))

    if record.get("bom"):
        files.append(_book(record, out_dir, "BOM", [(
            "BOM", "Bill of materials",
            ["Find", "P/N", "Description", "Qty", "UOM", "Material", "Spec", "Make/Buy", "Source"],
            [[b["find"], b["pn"], b["description"], b["qty"], b.get("uom", "EA"), b.get("material", ""), b.get("spec", ""),
              b.get("make_buy", ""), b.get("source", "")] for b in record["bom"]],
            [6, 18, 34, 6, 6, 18, 20, 9, 20], None)]))

    if record.get("design_reviews"):
        rows = []
        for d in record["design_reviews"]:
            for f in d["findings"]:
                rows.append([d["id"], d["gate"], d["date"], d["reviewer"], "yes" if d.get("signed") else "no", f["id"],
                             f["severity"], f.get("area", ""), f["text"], ", ".join(f.get("char_ids", [])), f["status"], f.get("resolution", "")])
        files.append(_book(record, out_dir, "Design-Review", [(
            "Design-Review", "Design review record (AS9100D 8.3.4)",
            ["Review", "Gate", "Date", "Reviewer", "Signed", "Finding", "Severity", "Area", "Finding text", "Char", "Status", "Resolution"],
            rows, [8, 5, 11, 14, 7, 8, 9, 12, 50, 8, 8, 40], None)]))

    if record.get("stackups"):
        summary, detail = [], []
        for s in record["stackups"]:
            res = analyze(s, chars)
            t = res["target"]
            summary.append([s["id"], s["name"], t["min"], t["max"], round(res["mean"], 5),
                            round(res["worst_case"]["min"], 5), round(res["worst_case"]["max"], 5),
                            round(res["rss"]["min"], 5), round(res["rss"]["max"], 5),
                            round(res["monte_carlo"]["sigma"], 5), f"{res['monte_carlo']['in_target']:.4%}",
                            res["method"], "PASS" if res["meets"] else "FAIL"])
            for c in res["contributors"]:
                detail.append([s["id"], c["name"], c.get("char_id", ""), "+" if c["sign"] > 0 else "-", c.get("nominal"),
                               c["lsl"], c["usl"], round((c["usl"] - c["lsl"]) / 2, 5), c.get("dist", "normal")])
        files.append(_book(record, out_dir, "Stack-ups", [
            ("Summary", "Tolerance stack-ups (worst case, RSS, Monte Carlo 3-sigma = half tolerance)",
             ["ID", "Stack", "Target min", "Target max", "Mean", "WC min", "WC max", "RSS min", "RSS max", "MC sigma",
              "MC in target", "Judged by", "Result"], summary, [6, 34, 9, 9, 9, 9, 9, 9, 9, 9, 10, 10, 7], None),
            ("Contributors", "Stack-up contributors",
             ["Stack", "Contributor", "Char", "Sign", "Nominal", "LSL", "USL", "Half tol", "Dist"],
             detail, [6, 34, 6, 5, 9, 9, 9, 9, 8], None)]))

    if record.get("dfmea"):
        files.append(_book(record, out_dir, "DFMEA", [(
            "DFMEA", "Design FMEA (AIAG-VDA 2019)", _FMEA_COLS, _fmea_rows(record["dfmea"]), _FMEA_W, 15)]))

    if record.get("process"):
        files.append(_book(record, out_dir, "Routing", [(
            "Routing", "Routing and process flow",
            ["Op", "Operation", "Work center", "Machine", "Setup / datums", "Makes", "Special process", "Instructions"],
            [[p["op"], p["name"], p.get("workcenter", ""), p.get("machine", ""), p.get("setup", ""),
              ", ".join(p.get("makes", [])), p.get("special_process", ""), p.get("instructions", "")] for p in record["process"]],
            [5, 26, 12, 16, 30, 20, 16, 40], None)]))

    if record.get("pfmea"):
        files.append(_book(record, out_dir, "PFMEA", [(
            "PFMEA", "Process FMEA (AIAG-VDA 2019)", _FMEA_COLS, _fmea_rows(record["pfmea"]), _FMEA_W, 15)]))

    if record.get("control_plan"):
        rows = []
        for r in record["control_plan"]:
            c = by_id.get(r["char_id"], {})
            rows.append([r.get("phase", "production"), r["op"], r["char_id"], c.get("feature", ""), c.get("class", ""),
                         _tol_text(c), r["method"], r.get("gauge", ""), r["sample_size"], r["frequency"],
                         r.get("control_method", ""), r["reaction_plan"]])
        files.append(_book(record, out_dir, "Control-Plan", [(
            "Control-Plan", "Control plan (AIAG Control Plan 2024 / AS9145)",
            ["Phase", "Op", "Char", "Characteristic", "Class", "Specification", "Evaluation method", "Gauge",
             "Sample size", "Frequency", "Control method", "Reaction plan"],
            rows, [10, 5, 6, 26, 6, 22, 22, 18, 9, 14, 22, 34], None)]))

    if record.get("inspection_plan"):
        rows = []
        for r in record["inspection_plan"]:
            c = by_id.get(r["char_id"], {})
            rows.append([r["char_id"], c.get("balloon", ""), c.get("feature", ""), c.get("class", ""), _tol_text(c),
                         r["stage"], r["method"], r.get("gauge", ""), r.get("sample", ""), r.get("notes", "")])
        files.append(_book(record, out_dir, "Inspection-Plan", [(
            "Inspection-Plan", "Inspection plan",
            ["Char", "Balloon", "Characteristic", "Class", "Specification", "Stage", "Method", "Gauge", "Sample", "Notes"],
            rows, [6, 7, 28, 6, 22, 10, 24, 18, 14, 30], None)]))

    if record.get("fai") or record.get("inspection_plan"):
        files.append(_fai(record, out_dir))

    if record.get("changes"):
        files.append(_book(record, out_dir, "ECN-Log", [(
            "ECN-Log", "Engineering change log (EIA-649C)",
            ["ECN", "Date", "From", "To", "Class", "Description", "Reason", "Char", "Affected documents", "FAI", "Approved by", "Status"],
            [[c["ecn"], c["date"], c["from_rev"], c["to_rev"], c.get("class", ""), c["description"], c["reason"],
              ", ".join(c.get("char_ids", [])), ", ".join(c.get("affected_docs", [])), c.get("fai_required", ""),
              c.get("approved_by", ""), c.get("status", "")] for c in record["changes"]],
            [9, 11, 5, 5, 6, 36, 30, 10, 30, 8, 14, 9], None)]))

    files.append(gate_report(record, out_dir))
    return files


def _fai(record, out_dir):
    """AS9102 Rev C Forms 1-3. Field layout follows the published form; check it against the standard you hold."""
    p = record["part"]
    fai = record.get("fai", {})
    form1 = [
        ["1. Part number", p.get("pn", p["id"])], ["2. Part name", p["name"]], ["3. Serial number", fai.get("serial", "")],
        ["4. FAI report number", fai.get("report_no", "")], ["5. Part revision level", p["rev"]],
        ["6. Drawing number", p.get("pn", "")], ["7. Drawing revision level", p["rev"]],
        ["8. Additional changes", ""], ["9. Manufacturing process reference", "Routing (see Routing workbook)"],
        ["10. Organization name", fai.get("organization", "")], ["11. Supplier code", p.get("cage", "")],
        ["12. P.O. number", ""], ["13. Detail / assembly", "assembly" if len(record.get("bom", [])) > 1 else "detail"],
        ["14. Full / partial FAI", fai.get("type", "full")], ["   Reason for partial FAI", fai.get("reason", "")],
        ["19. FAI complete / not complete", "complete" if fai.get("results") else "not complete"],
        ["20. Signature / date", ""], ["21. Reviewed by / date", ""], ["22. Customer approval / date", ""],
    ]
    form2 = [[m.get("item", ""), m.get("spec", ""), m.get("code", ""), m.get("supplier", ""), m.get("cert", ""), m.get("approved", "")]
             for m in fai.get("materials", []) + fai.get("special_processes", [])]
    form2 += [[t.get("item", ""), t.get("procedure", ""), "", "", t.get("result", ""), ""] for t in fai.get("functional_tests", [])]
    results = {r["char_id"]: r for r in fai.get("results", [])}
    form3 = []
    for c in record["characteristics"]:
        r = results.get(c["id"], {})
        form3.append([c.get("balloon") or c["id"][1:], c.get("drawing_zone", ""), c["class"] if c["class"] != "std" else "",
                      f"{c['feature']}: {_tol_text(c)}", r.get("result", ""), r.get("tooling", ""), r.get("ncr", "")])
    return _book(record, out_dir, "FAI", [
        ("FAI-Form1", "AS9102 Form 1: Part number accountability", ["Field", "Value"], form1, [36, 50], None),
        ("FAI-Form2", "AS9102 Form 2: Product accountability (raw material, special processes, functional testing)",
         ["Material / process / test", "Specification", "Code", "Supplier", "Certificate of conformance", "Customer approval"],
         form2, [26, 26, 10, 20, 24, 16], None),
        ("FAI-Form3", "AS9102 Form 3: Characteristic accountability, verification and compatibility evaluation",
         ["Char no.", "Ref. location", "Designator", "Requirement", "Results", "Designed / qualified tooling", "Nonconformance no."],
         form3, [8, 10, 10, 44, 14, 22, 16], None),
    ])


def gate_report(record, out_dir, gate=None):
    """Markdown report: the furthest gate that passes, issues for the next one, and KC/CC traceability."""
    p = record["part"]
    reached = None
    for g in GATES:
        if passed(check(record, g)):
            reached = g
        else:
            break
    if gate:
        target = gate
    elif reached is None:
        target = "G0"
    else:
        target = GATES[min(GATES.index(reached) + 1, len(GATES) - 1)]
    issues = check(record, target)
    lines = [f"# {p['id']} gate report (Rev - {p['rev']}, {dt.date.today().isoformat()})", ""]
    if p.get("mock"):
        lines += [f"**{MOCK_BANNER}**", ""]
    if p["export_control"].get("itar"):
        lines += [f"**{ITAR_WARNING}**", ""]
    lines += [f"Part: {p['name']}. P/N {p.get('pn', '-')}, NSN {p.get('nsn', '-')}.", "",
              f"Highest gate passed: **{reached or 'none'}**. Checked against: **{target}**.", ""]
    errs = [i for i in issues if i.level == "error"]
    warns = [i for i in issues if i.level == "warn"]
    lines += [f"## {target}: {len(errs)} errors, {len(warns)} warnings", ""]
    for i in errs + warns:
        lines.append(f"- {i.level.upper()} `{i.code}` ({i.gate}): {i.message}")
    if not issues:
        lines.append("No issues.")
    lines += ["", "## KC / CC traceability", "",
              "| Char | Feature | Class | Source | Req | DFMEA | Op | PFMEA | Control plan | Inspection | FAI |",
              "|---|---|---|---|---|---|---|---|---|---|---|"]
    for row in trace(record):
        c = row["char"]

        def cell(v):
            return ", ".join(str(x) for x in v) if v else "**missing**"

        lines.append(f"| {c['id']} | {c['feature']} | {c['class']} | {c['source']} | {cell(row['requirements'])} | {cell(row['dfmea'])} | "
                     f"{cell(row['ops'])} | {cell(row['pfmea'])} | {cell(row['control_plan'])} | {cell(row['inspection'])} | "
                     f"{'yes' if row['fai'] else '**missing**'} |")
    path = Path(out_dir) / f"{file_stem(record, 'Gate-Report')}.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path
