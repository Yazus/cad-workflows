import copy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engsys import record as rec
from engsys.ap import action_priority
from engsys.cli import main
from engsys.stackup import analyze
from engsys.workbooks import build

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "EX-0001" / "part-record.json"


@pytest.fixture
def r():
    return rec.load(EXAMPLE)


def codes(issues, level="error"):
    return {i.code for i in issues if i.level == level}


def test_example_passes_every_gate(r):
    assert rec.validate(r) == []
    for g in rec.GATES:
        assert rec.passed(rec.check(r, g)), [str(i) for i in rec.check(r, g)]


@pytest.mark.parametrize("s,o,d,ap", [
    (10, 10, 10, "H"), (9, 2, 2, "L"), (9, 3, 7, "H"), (9, 4, 1, "M"), (8, 4, 3, "M"),
    (7, 6, 1, "M"), (5, 8, 7, "H"), (5, 6, 1, "L"), (3, 8, 7, "M"), (3, 6, 10, "L"), (1, 10, 10, "L"), (10, 1, 10, "L"),
])
def test_action_priority(s, o, d, ap):
    assert action_priority(s, o, d) == ap


def test_action_priority_rejects_out_of_range():
    with pytest.raises(ValueError):
        action_priority(11, 1, 1)


def test_stackup_worst_case_and_rss(r):
    res = analyze(r["stackups"][0], r["characteristics"])
    assert res["worst_case"]["min"] == pytest.approx(0.0005)
    assert res["worst_case"]["max"] == pytest.approx(0.0015)
    assert res["rss"]["max"] - res["mean"] == pytest.approx((0.00025**2 * 2) ** 0.5)
    assert res["monte_carlo"]["mean"] == pytest.approx(0.001, abs=1e-5)
    assert res["meets"]


def test_stackup_failure_blocks_g2(r):
    bore = next(c for c in r["characteristics"] if c["id"] == "C4")
    bore["lsl"], bore["usl"] = 1.124, 1.1245  # the original bore DR1 caught: zero interference at worst case
    assert "stackup-fails" in codes(rec.check(r, "G2"))


def test_assumed_kc_blocks_g2(r):
    r["characteristics"][3]["source"] = "ASSUMED"
    assert "kc-assumed" in codes(rec.check(r, "G2"))


def test_high_ap_without_action_blocks_g2(r):
    line = next(l for l in r["dfmea"] if l["id"] == "D3")
    line.update(S=9, O=5, D=5)
    del line["actions"]
    assert "high-ap-no-action" in codes(rec.check(r, "G2"))


def test_open_major_finding_blocks_g2(r):
    r["design_reviews"][0]["findings"][0]["status"] = "open"
    assert "open-finding" in codes(rec.check(r, "G2"))


def test_orphan_requirement_blocks_g2(r):
    r["requirements"].append({"id": "R9", "type": "load", "text": "Unlinked", "source": "test"})
    assert "orphan-requirement" in codes(rec.check(r, "G2"))


def test_kc_missing_from_control_plan_blocks_g3(r):
    r["control_plan"] = [x for x in r["control_plan"] if x["char_id"] != "C9"]
    assert "kc-not-in-control-plan" in codes(rec.check(r, "G3"))


def test_char_not_made_blocks_g3(r):
    for p in r["process"]:
        p["makes"] = [c for c in p.get("makes", []) if c != "C13"]
    assert "char-not-made" in codes(rec.check(r, "G3"))


def test_pfmea_severity_below_linked_dfmea(r):
    r["pfmea"][0]["S"] = 5  # P1 links D1 (S8)
    assert "pfmea-s-below-dfmea" in codes(rec.check(r, "G3"))


def test_fai_reject_needs_ncr(r):
    res = next(x for x in r["fai"]["results"] if x["char_id"] == "C4")
    res["result"] = 1.1250
    assert "fai-reject-no-ncr" in codes(rec.check(r, "G4"))
    res["ncr"] = "NCR-001"
    assert "fai-reject-no-ncr" not in codes(rec.check(r, "G4"))


def test_ecn_rev_must_match(r):
    r["changes"].append({"ecn": "ECN-001", "date": "2026-10-21", "from_rev": "A", "to_rev": "B", "description": "Bore change",
                         "reason": "test", "char_ids": ["C4"], "affected_docs": ["drawing"], "status": "approved"})
    assert "rev-mismatch" in codes(rec.check(r, "G5"))
    r["part"]["rev"] = "B"
    assert rec.passed(rec.check(r, "G5"))


def test_mock_source_needs_mock_part(r):
    r["characteristics"][0]["source"] = "MOCK"
    assert "mock-in-real-record" in codes(rec.check(r, "G0"))


def test_schema_error_is_reported(r):
    del r["part"]["export_control"]
    assert "schema" in codes(rec.check(r, "G0"))


def test_build_writes_documents(r, tmp_path):
    files = build(r, tmp_path)
    names = {f.name for f in files}
    for doc in ("Characteristics", "DFMEA", "PFMEA", "Control-Plan", "Inspection-Plan", "FAI", "BOM", "Routing", "Stack-ups"):
        assert f"EX-0001_{doc} (Rev - A).xlsx" in names
    assert "EX-0001_Gate-Report (Rev - A).md" in names
    wb = load_workbook(tmp_path / "EX-0001_FAI (Rev - A).xlsx")
    assert wb.sheetnames == ["FAI-Form1", "FAI-Form2", "FAI-Form3"]


def test_mock_and_itar_marking(r, tmp_path):
    m = copy.deepcopy(r)
    m["part"]["mock"] = True
    m["part"]["export_control"]["itar"] = True
    files = build(m, tmp_path)
    assert all("MOCK" in f.name for f in files)
    ws = load_workbook(next(f for f in files if "DFMEA" in f.name)).active
    text = " ".join(str(c.value) for row in ws.iter_rows(max_row=5) for c in row if c.value)
    assert "MOCK-UP ONLY" in text and "ITAR" in text


def test_cli_check_and_init(tmp_path, capsys):
    assert main(["check", str(EXAMPLE), "--gate", "G4"]) == 0
    new = tmp_path / "rec.json"
    assert main(["init", str(new), "--id", "NEW-1", "--name", "New part"]) == 0
    assert rec.validate(json.loads(new.read_text())) == []
    assert main(["init", str(new), "--id", "NEW-1", "--name", "New part"]) == 2
    assert main(["check", str(new), "--gate", "G2"]) == 1
